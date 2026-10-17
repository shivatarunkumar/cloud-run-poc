"""Convert Excel (.xlsx) files to CSV so they go through the same validation and load path."""

import csv
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


def _cell_to_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, datetime):
        # Excel stores dates as datetimes; a midnight value is a plain date.
        return value.date().isoformat() if value.time() == time(0) else value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def xlsx_to_csv(src: Path, dst: Path) -> int:
    """Write the first worksheet of `src` to `dst` as UTF-8 CSV. Returns the number of lines written.

    The first non-empty row sets the width. Shorter rows (trailing NULLs) are padded to it; wider rows
    are kept as-is so the column-count check still catches them.
    """
    wb = load_workbook(src, read_only=True, data_only=True)
    try:
        ws = wb.worksheets[0]
        lines = 0
        width: int | None = None
        with dst.open("w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            for row in ws.iter_rows(values_only=True):
                cells = [_cell_to_text(v) for v in row]
                # Read-only mode can report empty cells beyond the data for formatted columns.
                while cells and cells[-1] == "":
                    cells.pop()
                if not cells:
                    continue
                if width is None:
                    width = len(cells)
                cells.extend([""] * (width - len(cells)))
                writer.writerow(cells)
                lines += 1
        return lines
    finally:
        wb.close()
